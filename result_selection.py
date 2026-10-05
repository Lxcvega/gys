import numpy as np
from typing import Dict, List, Optional, Tuple, Any
from dataclasses import dataclass
from data_model import (
    SupplierCapability, WastewaterType,
    SelectionResult, ChemicalRequirement, SupplierProposal,
    SupplierScoringResult, SupplierPoolMember
)
from nn_filter import SupplierFilterResult
from dynamic_game import GameResult, GameStrategy
@dataclass
class FinalSelectionCriteria:
    fuzzy_weight: float
    filter_weight: float
    nsga2_weight: float
    game_weight: float
    risk_tolerance: float
    budget_constraint: float
    quality_threshold: float

@dataclass
class ContractTerms:
    contract_duration: int
    payment_terms: str
    quality_guarantee: str
    delivery_requirements: str
    penalty_clauses: List[str]
    bonus_clauses: List[str]
    environmental_requirements: str

class FinalSelectionModule:
    def __init__(self):
        self.selection_criteria = FinalSelectionCriteria(
            fuzzy_weight=0.25,
            filter_weight=0.20,
            nsga2_weight=0.25,
            game_weight=0.30,
            risk_tolerance=0.7,
            budget_constraint=1.0,
            quality_threshold=0.6
        )
        
        self.winning_threshold = 0.65
        self.alternative_threshold = 0.55
        
        self.contract_templates = {
            WastewaterType.INDUSTRIAL: ContractTerms(
                contract_duration=24,
                payment_terms="分期付款，验收合格后支付80%，质保期后支付20%",
                quality_guarantee="重金属去除率≥95%，COD去除率≥85%",
                delivery_requirements="紧急订单24小时内响应，常规订单48小时内交付",
                penalty_clauses=[
                    "质量不达标扣除当批次货款的10%",
                    "延迟交付每日扣除合同金额的0.5%",
                    "环保不达标承担全部治理费用"
                ],
                bonus_clauses=[
                    "提前交付奖励合同金额的2%",
                    "质量超标奖励合同金额的5%"
                ],
                environmental_requirements="必须持有ISO14001认证，定期提交环保检测报告"
            ),
            WastewaterType.DOMESTIC: ContractTerms(
                contract_duration=12,
                payment_terms="月结30天",
                quality_guarantee="COD去除率≥80%，BOD去除率≥75%",
                delivery_requirements="常规订单72小时内交付",
                penalty_clauses=[
                    "质量不达标扣除当批次货款的8%",
                    "延迟交付每日扣除合同金额的0.3%"
                ],
                bonus_clauses=[
                    "连续3个月质量达标奖励合同金额的3%"
                ],
                environmental_requirements="符合国家生活污水处理标准"
            ),
            WastewaterType.MEDICAL: ContractTerms(
                contract_duration=18,
                payment_terms="预付30%，验收合格后支付60%，质保期后支付10%",
                quality_guarantee="病原体去除率≥99.9%，消毒效果100%",
                delivery_requirements="紧急订单12小时内响应，常规订单24小时内交付",
                penalty_clauses=[
                    "质量不达标扣除当批次货款的15%",
                    "延迟交付每日扣除合同金额的1%",
                    "安全事故承担全部责任"
                ],
                bonus_clauses=[
                    "零安全事故奖励合同金额的5%",
                    "提前交付奖励合同金额的3%"
                ],
                environmental_requirements="必须持有医疗废物处理资质，定期提交安全检测报告"
            ),
            WastewaterType.AGRICULTURAL: ContractTerms(
                contract_duration=12,
                payment_terms="季度结算",
                quality_guarantee="氮磷去除率≥70%",
                delivery_requirements="常规订单48小时内交付",
                penalty_clauses=[
                    "质量不达标扣除当批次货款的5%",
                    "延迟交付每日扣除合同金额的0.2%"
                ],
                bonus_clauses=[
                    "季节性需求满足率100%奖励合同金额的2%"
                ],
                environmental_requirements="符合农业面源污染治理标准"
            ),
            WastewaterType.CHEMICAL: ContractTerms(
                contract_duration=24,
                payment_terms="分期付款，验收合格后支付75%，质保期后支付25%",
                quality_guarantee="化学污染物去除率≥90%，pH调节精度±0.5",
                delivery_requirements="紧急订单8小时内响应，常规订单24小时内交付",
                penalty_clauses=[
                    "质量不达标扣除当批次货款的12%",
                    "延迟交付每日扣除合同金额的0.8%",
                    "环保事故承担全部治理费用及罚款"
                ],
                bonus_clauses=[
                    "技术创新奖励合同金额的4%",
                    "提前交付奖励合同金额的2%"
                ],
                environmental_requirements="必须持有危险化学品经营许可证，定期提交安全环保报告"
            )
        }

    def calculate_final_score(self, supplier_id: str,
                            fuzzy_scores: Dict[str, float],
                            filter_results: Dict[str, SupplierFilterResult],
                            nsga2_results: Dict[str, SupplierPoolMember],
                            game_results: Dict[str, GameResult],
                            scoring_results: Optional[Dict[str, SupplierScoringResult]] = None) -> float:
        fuzzy_score = fuzzy_scores.get(supplier_id, 0.0)
        
        filter_result = filter_results.get(supplier_id)
        filter_score = filter_result.overall_score if filter_result else 0.0
        
        nsga2_result = nsga2_results.get(supplier_id)
        nsga2_score = 1.0 / (nsga2_result.pareto_rank + 1) if nsga2_result else 0.0
        
        game_result = game_results.get(supplier_id)
        game_score = game_result.game_score if game_result else 0.0
        
        scoring_result = scoring_results.get(supplier_id) if scoring_results else None
        scoring_score = scoring_result.total_internal_score / 100 if scoring_result else 0.0

        w = self.selection_criteria
        final_score = (
            fuzzy_score * w.fuzzy_weight +
            filter_score * w.filter_weight +
            nsga2_score * w.nsga2_weight +
            game_score * w.game_weight +
            scoring_score * (1.0 - w.fuzzy_weight - w.filter_weight - w.nsga2_weight - w.game_weight)
        )
        
        return max(0, min(1, final_score))

    def assess_risk(self, supplier: SupplierCapability,
                   filter_result: Optional[SupplierFilterResult],
                   game_result: Optional[GameResult],
                   nsga2_result: Optional[SupplierPoolMember] = None) -> float:
        risk_factors = []

        risk_factors.append(1.0 - (supplier.safety_management_level / 5.0))
        risk_factors.append(1.0 - (supplier.delivery_capability / 5.0))
        risk_factors.append(1.0 - (supplier.emergency_supply_capability / 5.0))

        if filter_result:
            risk_factors.append(1.0 - filter_result.authenticity_score)

        if nsga2_result is not None:
            is_dummy = (
                nsga2_result.pareto_rank == 0
                and nsga2_result.crowding_distance == 0.5
                and all(v == 0.5 for v in nsga2_result.objectives.values())
                and nsga2_result.dominance_count == 0
            )
            if not is_dummy:
                nsga2_mitigation = 0.08 * (1.0 / (1.0 + nsga2_result.pareto_rank))
                risk_factors.append(-nsga2_mitigation)

        if game_result:
            game_risk_reduction = game_result.game_score * 0.12
            if game_result.strategy == GameStrategy.COOPERATIVE:
                game_risk_reduction += 0.05
            elif game_result.strategy == GameStrategy.MIXED:
                game_risk_reduction += 0.02
            risk_factors.append(-game_risk_reduction)

        overall_risk = float(np.mean(risk_factors)) if risk_factors else 0.5
        overall_risk = max(0.0, min(1.0, overall_risk))
        return overall_risk

    def calculate_expected_benefit(self, supplier: SupplierCapability,
                                   proposal: SupplierProposal,
                                   requirement: ChemicalRequirement) -> float:
        quality_benefit = (supplier.effective_ingredient_content / 100.0) * 0.30
        cost_benefit = (1.0 - min(1.0, proposal.unit_price / max(1e-6, requirement.budget_constraint))) * 0.25
        delivery_benefit = (supplier.delivery_timeliness / 100.0) * 0.20
        technical_benefit = (supplier.delivery_capability / 5.0) * 0.15
        environmental_benefit = (supplier.safety_management_level / 5.0) * 0.10
        
        total_benefit = quality_benefit + cost_benefit + delivery_benefit + technical_benefit + environmental_benefit
        return max(0.0, total_benefit)

    def determine_selection_status(self, final_score: float, risk_level: float) -> Tuple[bool, str]:
        risk_adjusted_score = final_score * (1.0 - risk_level * 0.3)
        
        if risk_adjusted_score >= self.winning_threshold:
            return True, "推荐"
        elif risk_adjusted_score >= self.alternative_threshold:
            return False, "预备推荐"
        else:
            return False, "未推荐"

    def _get_internal_score(self, selection_result: SelectionResult) -> float:
        return selection_result.contract_terms.get('_internal_score', 0.0)

    def generate_contract_terms(self, supplier: SupplierCapability,
                              wastewater_type: WastewaterType,
                              proposal: SupplierProposal) -> Dict[str, Any]:
        template = self.contract_templates.get(wastewater_type)
        
        if not template:
            template = self.contract_templates[WastewaterType.DOMESTIC]
        
        contract_terms = {
            'supplier_id': supplier.supplier_id,
            'supplier_name': supplier.supplier_name,
            'contract_duration': template.contract_duration,
            'payment_terms': template.payment_terms,
            'quality_guarantee': template.quality_guarantee,
            'delivery_requirements': template.delivery_requirements,
            'penalty_clauses': template.penalty_clauses.copy(),
            'bonus_clauses': template.bonus_clauses.copy(),
            'environmental_requirements': template.environmental_requirements,
            'unit_price': proposal.unit_price,
            'offered_quantity': proposal.offered_quantity,
            'delivery_time': proposal.delivery_time,
            'quality_certifications': proposal.quality_certification,
            'environmental_certifications': proposal.environmental_certification
        }
        
        if supplier.emergency_supply_capability >= 4:
            contract_terms['bonus_clauses'].append("应急供应能力优秀奖励合同金额的3%")
        
        if supplier.safety_management_level >= 4:
            contract_terms['bonus_clauses'].append("安全管理水平优秀奖励合同金额的2%")
        
        return contract_terms

    def finalize_selection(self, suppliers: List[SupplierCapability],
                          proposals: Dict[str, SupplierProposal],
                          fuzzy_scores: Dict[str, float],
                          filter_results: Dict[str, SupplierFilterResult],
                          nsga2_results: Dict[str, SupplierPoolMember],
                          game_results: Dict[str, GameResult],
                          requirements: List[ChemicalRequirement],
                          wastewater_type: WastewaterType,
                          scoring_results: Optional[Dict[str, SupplierScoringResult]] = None) -> List[SelectionResult]:
        selection_results = []
        
        for supplier in suppliers:
            supplier_id = supplier.supplier_id
            proposal = proposals.get(supplier_id)
            
            if not proposal:
                continue
            
            final_score = self.calculate_final_score(
                supplier_id, fuzzy_scores, filter_results,
                nsga2_results, game_results, scoring_results
            )
            
            filter_result = filter_results.get(supplier_id)
            game_result = game_results.get(supplier_id)
            nsga2_result = nsga2_results.get(supplier_id)
            
            risk_level = self.assess_risk(supplier, filter_result, game_result, nsga2_result)
            
            requirement = next((req for req in requirements 
                              if req.chemical_type == proposal.chemical_type), None)
            
            if not requirement:
                continue
            
            expected_benefit = self.calculate_expected_benefit(supplier, proposal, requirement)
            
            is_selected, selection_status = self.determine_selection_status(final_score, risk_level)
            
            scoring_result = scoring_results.get(supplier_id) if scoring_results else None
            internal_score = scoring_result.total_internal_score if scoring_result else 0.0
            
            contract_terms = self.generate_contract_terms(supplier, wastewater_type, proposal)
            contract_terms['_internal_score'] = internal_score
            
            selection_reason = self._generate_selection_reason(
                final_score, risk_level, is_selected,
                filter_result, game_result, scoring_result
            )
            
            selection_results.append(SelectionResult(
                supplier_id=supplier_id,
                supplier_name=supplier.supplier_name,
                final_score=final_score,
                is_selected=is_selected,
                selection_reason=selection_reason,
                risk_assessment=risk_level,
                expected_benefit=expected_benefit,
                contract_terms=contract_terms
            ))
        
        sorted_results = sorted(selection_results, key=lambda x: self._get_internal_score(x), reverse=True)
        
        total_suppliers = len(sorted_results)
        
        recommended_count = min(3, total_suppliers)
        for i in range(recommended_count):
            if not sorted_results[i].is_selected:
                sorted_results[i].is_selected = True
                if "暂不符合推荐条件" in sorted_results[i].selection_reason:
                    sorted_results[i].selection_reason = sorted_results[i].selection_reason.replace(
                        "暂不符合推荐条件", "符合推荐条件"
                    )
                elif "未推荐" in sorted_results[i].selection_reason:
                    sorted_results[i].selection_reason = sorted_results[i].selection_reason.replace(
                        "未推荐", "慎重选择"
                    )
        
        alternative_count = min(6, total_suppliers - recommended_count)
        for i in range(recommended_count, recommended_count + alternative_count):
            if "未推荐" in sorted_results[i].selection_reason:
                sorted_results[i].selection_reason = sorted_results[i].selection_reason.replace("未推荐", "预备推荐")
        
        selected_count = sum(1 for r in sorted_results if r.is_selected)
        if selected_count > 9:
            kept = 0
            for r in sorted_results:
                if r.is_selected:
                    kept += 1
                    if kept > 9:
                        r.is_selected = False
                        if "推荐" in r.selection_reason:
                            r.selection_reason = r.selection_reason.replace("推荐", "未推荐")
                        if "取消推荐" not in r.selection_reason:
                            r.selection_reason = r.selection_reason + "；因配额限制取消推荐"
        
        return sorted_results

    def _generate_selection_reason(self, final_score: float, risk_level: float,
                                   is_selected: bool,
                                   filter_result: Optional[SupplierFilterResult],
                                   game_result: Optional[GameResult],
                                   scoring_result: Optional[SupplierScoringResult] = None) -> str:
        reasons = []
        
        if final_score >= self.winning_threshold:
            reasons.append(f"综合评分优秀({final_score:.3f})")
        elif final_score >= self.alternative_threshold:
            reasons.append(f"综合评分良好({final_score:.3f})")
        else:
            reasons.append(f"综合评分一般({final_score:.3f})")
        
        if risk_level < 0.3:
            reasons.append("风险水平低")
        elif risk_level < 0.5:
            reasons.append("风险水平中等")
        else:
            reasons.append("风险水平较高")
        
        if filter_result and filter_result.authenticity_score > 0.8:
            reasons.append("数据真实性高")
        
        if game_result and game_result.strategy.value == "合作策略":
            reasons.append("合作态度良好")
        
        if scoring_result:
            if scoring_result.total_internal_score >= 80:
                reasons.append(f"内部评分优秀({scoring_result.total_internal_score:.1f})")
            elif scoring_result.total_internal_score >= 60:
                reasons.append(f"内部评分良好({scoring_result.total_internal_score:.1f})")
            else:
                reasons.append(f"内部评分一般({scoring_result.total_internal_score:.1f})")
        
        if is_selected:
            return ";".join(reasons) + "；符合推荐条件"
        else:
            return ";".join(reasons) + "；暂不符合推荐条件"

    def generate_selection_report(self, selection_results: List[SelectionResult],
                                 wastewater_type: WastewaterType) -> str:
        report = "供应商最终选择报告\n"
        report += "=" * 60 + "\n\n"
        
        total_suppliers = len(selection_results)
        
        recommended_count = min(3, total_suppliers)
        recommended_suppliers = selection_results[:recommended_count]
        
        alternative_count = min(6, total_suppliers - recommended_count)
        alternative_suppliers = selection_results[recommended_count:recommended_count + alternative_count] if alternative_count > 0 else []
        
        report += f"污水类型: {wastewater_type.value}\n"
        report += f"参与供应商: {total_suppliers} 家\n"
        report += f"推荐供应商: {len(recommended_suppliers)} 家\n"
        report += f"预备推荐供应商: {len(alternative_suppliers)} 家\n\n"
        
        if recommended_suppliers:
            report += "推荐供应商:\n"
            report += "-" * 60 + "\n"
            for i, result in enumerate(recommended_suppliers, 1):
                status = "推荐但是慎重选择" if "慎重选择" in result.selection_reason else "推荐"
                report += f"{i}. {result.supplier_name} (ID: {result.supplier_id}) - {status}\n"
                report += f"   内部评分: {self._get_internal_score(result):.1f}/100\n"
                report += f"   综合评分: {result.final_score:.4f} (仅供参考)\n"
                report += f"   风险评估: {result.risk_assessment:.4f}\n"
                report += f"   预期收益: {result.expected_benefit:.4f}\n"
                report += f"   选择原因: {result.selection_reason}\n"
                report += f"   合同期限: {result.contract_terms['contract_duration']} 个月\n"
                report += f"   单价: {result.contract_terms['unit_price']:.2f} 元\n"
                report += f"   交付时间: {result.contract_terms['delivery_time']} 天\n\n"
        
        if alternative_suppliers:
            report += "预备推荐供应商:\n"
            report += "-" * 60 + "\n"
            for i, result in enumerate(alternative_suppliers, 1):
                report += f"{i}. {result.supplier_name} (ID: {result.supplier_id})\n"
                report += f"   内部评分: {self._get_internal_score(result):.1f}/100\n"
                report += f"   综合评分: {result.final_score:.4f} (仅供参考)\n"
                report += f"   风险评估: {result.risk_assessment:.4f}\n"
                report += f"   选择原因: {result.selection_reason}\n\n"
        
        return report

    def export_selection_results(self, selection_results: List[SelectionResult],
                                output_file: str = "selection_results.txt") -> bool:
        try:
            total_suppliers = len(selection_results)
            recommended_count = min(3, total_suppliers)
            alternative_count = min(6, total_suppliers - recommended_count)
            export_count = recommended_count + alternative_count
            export_results = selection_results[:export_count]
            
            with open(output_file, 'w', encoding='utf-8') as f:
                for result in export_results:
                    f.write(f"供应商ID: {result.supplier_id}\n")
                    f.write(f"供应商名称: {result.supplier_name}\n")
                    f.write(f"内部评分: {self._get_internal_score(result):.1f}/100\n")
                    f.write(f"综合评分: {result.final_score:.4f} (仅供参考)\n")
                    f.write(f"是否推荐: {'是' if result.is_selected else '否'}\n")
                    f.write(f"选择原因: {result.selection_reason}\n")
                    f.write(f"风险评估: {result.risk_assessment:.4f}\n")
                    f.write(f"预期收益: {result.expected_benefit:.4f}\n")
                    f.write("合同条款:\n")
                    for key, value in result.contract_terms.items():
                        f.write(f"  {key}: {value}\n")
                    f.write("\n" + "="*60 + "\n\n")
            return True
        except Exception as e:
            print(f"导出结果失败: {e}")
            return False

    def export_selection_results_excel(self, selection_results: List[SelectionResult],
                                       wastewater_type: WastewaterType,
                                       output_file: str = "selection_results.xlsx") -> bool:
        """将最终选择结果导出为 Excel 表格"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            total_suppliers = len(selection_results)
            recommended_count = min(3, total_suppliers)
            alternative_count = min(6, total_suppliers - recommended_count)
            export_count = recommended_count + alternative_count
            export_results = selection_results[:export_count]

            def _txt(v):
                if isinstance(v, (list, tuple)):
                    return "；".join(str(x) for x in v)
                return "" if v is None else str(v)

            def _fmt_num(v, digits=2):
                try:
                    return round(float(v), digits)
                except (TypeError, ValueError):
                    return v if v is not None else ""

            wb = Workbook()
            ws = wb.active
            ws.title = "最终选择结果"

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
            title_font = Font(bold=True, size=14)
            center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            left = Alignment(horizontal="left", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin")
            )

            ws.merge_cells("A1:L1")
            ws["A1"] = "污水供应商最终选择结果"
            ws["A1"].font = title_font
            ws["A1"].alignment = center

            ws.merge_cells("A2:L2")
            ws["A2"] = (f"污水类型：{wastewater_type.value}    参与供应商：{total_suppliers} 家    "
                        f"推荐：{recommended_count} 家    预备推荐：{alternative_count} 家")
            ws["A2"].alignment = left

            headers = ["排名", "供应商编号", "供应商名称", "内部评分(100)", "综合评分", "推荐状态",
                       "风险评估", "预期收益", "合同期限(月)", "单价(元/吨)", "交付时间(天)", "选择原因"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = thin_border

            for i, result in enumerate(export_results, 1):
                row = 4 + i
                if i <= recommended_count:
                    status = "推荐（慎重选择）" if "慎重选择" in result.selection_reason else "推荐"
                else:
                    status = "预备推荐"
                ct = result.contract_terms
                values = [
                    i,
                    result.supplier_id,
                    result.supplier_name,
                    _fmt_num(self._get_internal_score(result), 1),
                    _fmt_num(result.final_score, 4),
                    status,
                    _fmt_num(result.risk_assessment, 4),
                    _fmt_num(result.expected_benefit, 4),
                    ct.get('contract_duration', ''),
                    _fmt_num(ct.get('unit_price')),
                    ct.get('delivery_time', ''),
                    result.selection_reason,
                ]
                for col_idx, value in enumerate(values, 1):
                    cell = ws.cell(row=row, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = center if col_idx != 12 else left

            widths = [6, 12, 16, 14, 10, 16, 10, 10, 14, 12, 14, 50]
            for col_idx, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            ws.freeze_panes = "A5"

            ws2 = wb.create_sheet("合同条款")
            ws2.merge_cells("A1:I1")
            ws2["A1"] = "推荐供应商合同条款"
            ws2["A1"].font = title_font
            ws2["A1"].alignment = center

            ct_headers = ["排名", "供应商编号", "供应商名称", "付款方式", "质量保证",
                          "交付要求", "违约条款", "奖励条款", "环保要求"]
            for col_idx, header in enumerate(ct_headers, 1):
                cell = ws2.cell(row=3, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = thin_border

            for i, result in enumerate(export_results, 1):
                row = 3 + i
                ct = result.contract_terms
                ct_values = [
                    i,
                    result.supplier_id,
                    result.supplier_name,
                    _txt(ct.get('payment_terms')),
                    _txt(ct.get('quality_guarantee')),
                    _txt(ct.get('delivery_requirements')),
                    _txt(ct.get('penalty_clauses')),
                    _txt(ct.get('bonus_clauses')),
                    _txt(ct.get('environmental_requirements')),
                ]
                for col_idx, value in enumerate(ct_values, 1):
                    cell = ws2.cell(row=row, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = center if col_idx <= 3 else left

            ct_widths = [6, 12, 16, 20, 24, 24, 40, 40, 30]
            for col_idx, width in enumerate(ct_widths, 1):
                ws2.column_dimensions[get_column_letter(col_idx)].width = width
            ws2.freeze_panes = "A4"

            wb.save(output_file)
            print(f"Excel结果已导出: {output_file}")
            return True
        except Exception as e:
            print(f"导出Excel结果失败: {e}")
            return False

    def export_selection_results_word(self, selection_results: List[SelectionResult],
                                      wastewater_type: WastewaterType,
                                      output_file: str = "selection_results.docx") -> bool:
        """将最终选择结果导出为 Word 文档"""
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            total_suppliers = len(selection_results)
            recommended_count = min(3, total_suppliers)
            alternative_count = min(6, total_suppliers - recommended_count)
            export_count = recommended_count + alternative_count
            export_results = selection_results[:export_count]

            def _txt(v):
                if isinstance(v, (list, tuple)):
                    return "；".join(str(x) for x in v)
                return "" if v is None else str(v)

            def _fmt_num(v, digits=2):
                try:
                    return f"{float(v):.{digits}f}"
                except (TypeError, ValueError):
                    return str(v) if v is not None else ""

            doc = Document()
            title = doc.add_heading("污水供应商最终选择报告", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            meta.add_run(
                f"污水类型：{wastewater_type.value}    参与供应商：{total_suppliers} 家    "
                f"推荐：{recommended_count} 家    预备推荐：{alternative_count} 家"
            ).font.size = Pt(10.5)

            doc.add_heading("一、供应商评分排序结果", level=1)
            headers = ["排名", "供应商编号", "供应商名称", "内部评分", "综合评分", "推荐状态",
                       "风险评估", "预期收益", "选择原因"]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for idx, header in enumerate(headers):
                hdr_cells[idx].text = header
                for p in hdr_cells[idx].paragraphs:
                    for r in p.runs:
                        r.font.bold = True

            for i, result in enumerate(export_results, 1):
                if i <= recommended_count:
                    status = "推荐（慎重选择）" if "慎重选择" in result.selection_reason else "推荐"
                else:
                    status = "预备推荐"
                row_values = [
                    str(i),
                    result.supplier_id,
                    result.supplier_name,
                    _fmt_num(self._get_internal_score(result), 1),
                    _fmt_num(result.final_score, 4),
                    status,
                    _fmt_num(result.risk_assessment, 4),
                    _fmt_num(result.expected_benefit, 4),
                    result.selection_reason,
                ]
                cells = table.add_row().cells
                for idx, value in enumerate(row_values):
                    cells[idx].text = value

            doc.add_heading("二、推荐供应商合同条款", level=1)
            for i, result in enumerate(export_results, 1):
                status = "推荐" if i <= recommended_count else "预备推荐"
                doc.add_heading(f"{i}. {result.supplier_name}（{result.supplier_id}）— {status}", level=2)
                ct = result.contract_terms
                fields = [
                    ("合同期限", f"{_fmt_num(ct.get('contract_duration'), 0)} 个月"),
                    ("单价", f"{_fmt_num(ct.get('unit_price'))} 元/吨"),
                    ("交付时间", f"{_fmt_num(ct.get('delivery_time'), 0)} 天"),
                    ("付款方式", _txt(ct.get('payment_terms'))),
                    ("质量保证", _txt(ct.get('quality_guarantee'))),
                    ("交付要求", _txt(ct.get('delivery_requirements'))),
                    ("违约条款", _txt(ct.get('penalty_clauses'))),
                    ("奖励条款", _txt(ct.get('bonus_clauses'))),
                    ("环保要求", _txt(ct.get('environmental_requirements'))),
                ]
                for label, value in fields:
                    p = doc.add_paragraph()
                    p.add_run(f"{label}：").bold = True
                    p.add_run(value)

            doc.save(output_file)
            print(f"Word结果已导出: {output_file}")
            return True
        except Exception as e:
            print(f"导出Word结果失败: {e}")
            return False