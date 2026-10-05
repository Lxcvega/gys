"""Release regressions: run with python -m unittest discover -s tests -v.

HDP orchestration is isolated from optional ML dependencies; spreadsheet and
core scoring/export checks execute the real implementation and real files.
"""
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def supplier_row(sid="S1", **changes):
    row = dict(zip(
        ['供应商编号', '供应商名称', '药剂单价(元/吨)', '运输成本(元/吨)',
         '付款周期(天)', '药剂有效成分含量(%)', '供货及时性(%)',
         '库存能力(吨)', '交付能力(1-5级)', '应急供应能力(1-5级)',
         '安全管理水平(1-5级)', '处理污泥效果达标率(%)'],
        [sid, 'Synthetic ' + str(sid), 1500, 150, 30, 95, 96, 5000, 4, 3, 4, 90]))
    row.update(changes)
    return row


class RootCliTests(unittest.TestCase):
    def setUp(self):
        self.root = load_module('release_root', 'main.py')

    def test_hdp_flags_preserved_and_argv_restored(self):
        options = ['--mode', 'full', '--data', 'file with spaces.xlsx',
                   '--historical', 'labels.xlsx', '--window', '8', '--epochs', '3']
        original = sys.argv
        seen = []
        target = types.ModuleType('HDP_DS.main')
        target.main = lambda: seen.append(list(sys.argv)) or 17
        with patch.dict(sys.modules, {'HDP_DS.main': target}):
            self.assertEqual(self.root.main(['--hdp', *options]), 17)
        self.assertEqual(seen, [[original[0], *options]])
        self.assertIs(sys.argv, original)

    def test_argv_restored_after_hdp_failure(self):
        original = sys.argv
        target = types.ModuleType('HDP_DS.main')
        target.main = Mock(side_effect=RuntimeError('training failed'))
        with patch.dict(sys.modules, {'HDP_DS.main': target}):
            with self.assertRaisesRegex(RuntimeError, 'training failed'):
                self.root.main(['--hdp', '--mode', 'full'])
        self.assertIs(sys.argv, original)

    def test_help_in_subprocess_does_not_import_optional_ml_or_prompt(self):
        result = subprocess.run([sys.executable, str(ROOT / 'main.py'), '--help'],
                                capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn(b'--hdp', result.stdout)
        self.assertNotIn(b'Traceback', result.stderr)

    def test_mixed_modes_rejected_before_dispatch(self):
        with contextlib.redirect_stderr(io.StringIO()) as errors:
            with self.assertRaises(SystemExit) as exit_info:
                self.root.main(['--hdp', '--legacy'])
        self.assertEqual(exit_info.exception.code, 2)
        self.assertIn('--legacy', errors.getvalue())

    def test_legacy_rejects_hdp_options(self):
        with contextlib.redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as exit_info:
                self.root.main(['--legacy', '--epochs', '2'])
        self.assertEqual(exit_info.exception.code, 2)

    def test_default_and_explicit_legacy_dispatch(self):
        target = types.ModuleType('legacy_main')
        target.main = Mock(return_value='legacy')
        with patch.dict(sys.modules, {'legacy_main': target}):
            self.assertEqual(self.root.main([]), 'legacy')
            self.assertEqual(self.root.main(['--legacy']), 'legacy')
        self.assertEqual(target.main.call_count, 2)


class SpreadsheetTests(unittest.TestCase):
    def setUp(self):
        from excel_data_reader import ExcelDataReader
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'suppliers.xlsx'
        self.reader = ExcelDataReader(str(self.path))

    def read(self, rows):
        pd.DataFrame(rows).to_excel(self.path, index=False)
        with contextlib.redirect_stdout(io.StringIO()) as output:
            result = self.reader.read_excel_data()
        return result, output.getvalue()

    def test_failure_clears_prior_data_and_summary(self):
        for failure in ('columns', 'missing', 'corrupt', 'empty'):
            with self.subTest(failure=failure):
                self.assertEqual(len(self.read([supplier_row()])[0]), 1)
                if failure == 'columns':
                    pd.DataFrame({'供应商编号': ['S2']}).to_excel(self.path, index=False)
                elif failure == 'missing':
                    self.path.unlink()
                elif failure == 'corrupt':
                    self.path.write_bytes(b'not an Excel file')
                else:
                    pd.DataFrame(columns=supplier_row().keys()).to_excel(self.path, index=False)
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(self.reader.read_excel_data(), [])
                self.assertEqual(self.reader.supplier_data, [])
                self.assertEqual(self.reader.get_all_suppliers(), [])
                self.assertEqual(self.reader.generate_data_summary(), {})

    def test_missing_and_normalized_duplicate_ids_all_excluded(self):
        rows = [supplier_row(' duplicate '), supplier_row('duplicate'),
                supplier_row(7), supplier_row('7'), supplier_row(None),
                supplier_row('  '), supplier_row('valid')]
        result, diagnostic = self.read(rows)
        self.assertEqual([r.supplier_id for r in result], ['valid'])
        self.assertIn('重复', diagnostic)
        self.assertIn('第', diagnostic)

    def test_numeric_ids_with_missing_cells_cannot_become_nan_ids(self):
        result, diagnostic = self.read([supplier_row(12), supplier_row(None),
                                       supplier_row(np.nan), supplier_row('')])
        self.assertEqual([r.supplier_id for r in result], ['12'])
        self.assertIn('编号不能为空', diagnostic)

    def test_invalid_rows_excluded_before_imputation(self):
        invalid = [('药剂单价(元/吨)', -1), ('运输成本(元/吨)', 'bad'),
                   ('付款周期(天)', -1), ('库存能力(吨)', np.inf),
                   ('药剂有效成分含量(%)', 101), ('供货及时性(%)', -0.1),
                   ('处理污泥效果达标率(%)', 101), ('交付能力(1-5级)', 6),
                   ('应急供应能力(1-5级)', 0), ('安全管理水平(1-5级)', 2.5),
                   ('药剂单价(元/吨)', 'nan')]
        rows = [supplier_row('good'), supplier_row('blank', **{'药剂单价(元/吨)': None})]
        rows += [supplier_row('bad' + str(i), **{column: value})
                 for i, (column, value) in enumerate(invalid)]
        result, diagnostic = self.read(rows)
        self.assertEqual([r.supplier_id for r in result], ['good', 'blank'])
        self.assertEqual(result[1].material_unit_price, 1500)
        self.assertEqual(diagnostic.count('数据无效，已排除'), len(invalid))

    def test_blank_grades_and_entire_numeric_column_are_imputed(self):
        rows = [supplier_row('a', **{'交付能力(1-5级)': 3, '运输成本(元/吨)': None}),
                supplier_row('b', **{'交付能力(1-5级)': 4, '运输成本(元/吨)': None}),
                supplier_row('c', **{'交付能力(1-5级)': None, '运输成本(元/吨)': None})]
        result, diagnostic = self.read(rows)
        self.assertEqual(len(result), 3)
        self.assertEqual(result[2].delivery_capability, 4)
        self.assertTrue(all(row.transport_cost == 150 for row in result))
        self.assertIn('填充', diagnostic)


class HDPControlFlowTests(unittest.TestCase):
    def setUp(self):
        # Load the actual HDP entrypoint without importing torch/model classes.
        package = types.ModuleType('HDP_DS')
        for name in ('HDPDataLoader HDPDataset HDPModel HDPModelConfig HDPNSGA2 '
                     'NSGA2Config TOPSISDecision TOPSISConfig EvaluationMetrics '
                     'ExperimentComparator AblationConfig AblationRunner FeedbackLoop '
                     'FeedbackConfig').split():
            setattr(package, name, Mock(name=name))
        with patch.dict(sys.modules, {'HDP_DS': package}):
            self.hdp = load_module('release_hdp', 'HDP_DS/main.py')
        self.hdp._run_demo = Mock(name='demo')

    def call_main(self, options):
        with patch.object(sys, 'argv', ['hdp', *options]), \
                contextlib.redirect_stdout(io.StringIO()), \
                contextlib.redirect_stderr(io.StringIO()):
            return self.hdp.main()

    def test_hdp_help_exits_without_loading_or_training(self):
        with self.assertRaises(SystemExit) as info:
            self.call_main(['--help'])
        self.assertEqual(info.exception.code, 0)
        self.hdp.HDPDataLoader.assert_not_called()
        self.hdp.HDPModel.assert_not_called()
        self.hdp._run_demo.assert_not_called()

    def test_root_hdp_help_reaches_hdp_parser_without_work(self):
        root = load_module('release_root_help', 'main.py')
        original = sys.argv
        with patch.dict(sys.modules, {'HDP_DS.main': self.hdp}), \
                contextlib.redirect_stdout(io.StringIO()) as output:
            with self.assertRaises(SystemExit) as info:
                root.main(['--hdp', '--help'])
        self.assertEqual(info.exception.code, 0)
        self.assertIn('--window', output.getvalue())
        self.assertIs(sys.argv, original)
        self.hdp.HDPModel.assert_not_called()
        self.hdp.HDPDataLoader.assert_not_called()

    def test_only_explicit_demo_runs_demo(self):
        self.call_main(['--mode', 'demo'])
        self.hdp._run_demo.assert_called_once_with()
        self.hdp.HDPDataLoader.assert_not_called()

    def test_all_real_modes_fail_missing_file_without_demo(self):
        with tempfile.TemporaryDirectory() as temp:
            missing = str(Path(temp) / 'missing.xlsx')
            for mode in ('full', 'ablation', 'comparison', 'feedback'):
                with self.subTest(mode=mode), self.assertRaises(SystemExit) as info:
                    self.call_main(['--mode', mode, '--data', missing])
                self.assertEqual(info.exception.code, 2)
        self.hdp._run_demo.assert_not_called()
        self.hdp.HDPModel.assert_not_called()

    def test_bad_empty_nonfinite_and_wrong_shape_data_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'input.xlsx'
            path.touch()
            loader = self.hdp.HDPDataLoader.return_value
            loader.load_from_excel.return_value = False
            with self.assertRaisesRegex(ValueError, '加载失败'):
                self.hdp._load_dataset(str(path), 2)
            loader.load_from_excel.return_value = True
            dataset = loader.build_dataset.return_value = Mock()
            dataset.__len__ = Mock(return_value=0)
            with self.assertRaisesRegex(ValueError, '滚动窗口'):
                self.hdp._load_dataset(str(path), 2)
            dataset.__len__.return_value = 3
            for arrays in ((np.full((3, 6), np.inf), np.zeros((3, 2, 6)), np.zeros((3, 5))),
                           (np.zeros((3, 5)), np.zeros((3, 2, 6)), np.zeros((3, 5)))):
                dataset.to_numpy.return_value = arrays
                with self.assertRaisesRegex(ValueError, '形状或数值无效'):
                    self.hdp._load_dataset(str(path), 2)
        self.hdp._run_demo.assert_not_called()

    def test_full_dispatch_loads_only_once_and_guards_empty_splits(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'input.xlsx'
            path.touch()
            loader = self.hdp.HDPDataLoader.return_value
            loader.load_from_excel.return_value = True
            dataset = loader.build_dataset.return_value = Mock()
            dataset.__len__ = Mock(return_value=3)
            dataset.to_numpy.return_value = (np.zeros((3, 6)), np.zeros((3, 2, 6)), np.zeros((3, 5)))
            empty = Mock()
            empty.__len__ = Mock(return_value=0)
            loader.train_val_test_split.return_value = (dataset, dataset, empty)
            with self.assertRaises(SystemExit) as info:
                self.call_main(['--data', str(path), '--window', '2', '--epochs', '1'])
            self.assertEqual(info.exception.code, 2)
            loader.load_from_excel.assert_called_once_with(str(path))
            loader.build_dataset.assert_called_once_with()
            self.hdp.HDPModel.assert_not_called()

    def test_nonpositive_parameters_rejected_without_work(self):
        for flag in ('--epochs', '--window'):
            with self.subTest(flag=flag), self.assertRaises(SystemExit) as info:
                self.call_main([flag, '0', '--mode', 'demo'])
            self.assertEqual(info.exception.code, 2)
        self.hdp._run_demo.assert_not_called()


class CoreIntegrationTests(unittest.TestCase):
    def test_historical_real_score_labels_are_unchanged(self):
        import legacy_main
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'historical.xlsx'
            labels = [0.0, 31.125, 100.0]
            pd.DataFrame([supplier_row(str(i), **{'实际评分': score})
                          for i, score in enumerate(labels)]).to_excel(path, index=False)
            system = legacy_main.WastewaterSupplierSelectionSystem()
            with patch.object(system.nn_filter_module, 'add_regression_training_data') as add, \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertTrue(system.load_historical_training_data(str(path)))
            self.assertEqual([call.args[1] for call in add.call_args_list], labels)

    def test_legacy_default_paths_are_project_relative_from_foreign_cwd(self):
        import legacy_main
        original = os.getcwd()
        with tempfile.TemporaryDirectory() as temp:
            try:
                os.chdir(temp)
                with patch('builtins.input', return_value=''), \
                        patch.object(legacy_main.WastewaterSupplierSelectionSystem,
                                     'run_complete_workflow', return_value=[]) as run, \
                        contextlib.redirect_stdout(io.StringIO()):
                    legacy_main.main()
                self.assertEqual(Path(run.call_args.args[3]), ROOT / 'data/raw/供应商数据.xlsx')
                self.assertEqual(Path(run.call_args.args[4]), ROOT / 'data/raw/历史训练数据_模板.xlsx')
            finally:
                os.chdir(original)

    def test_real_reader_scoring_selection_and_exports(self):
        from docx import Document
        from openpyxl import load_workbook
        from excel_data_reader import ExcelDataReader
        from supplier_scoring import SupplierScoringModule
        from result_selection import FinalSelectionModule
        from data_model import ChemicalType, ChemicalRequirement, WastewaterType
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'suppliers.xlsx'
            pd.DataFrame([supplier_row('A'), supplier_row('B', **{
                '药剂单价(元/吨)': 2000, '交付能力(1-5级)': 3})]).to_excel(path, index=False)
            reader = ExcelDataReader(str(path))
            with contextlib.redirect_stdout(io.StringIO()):
                raw = reader.read_excel_data()
            scoring = SupplierScoringModule().batch_calculate_scores(
                [(reader.convert_to_supplier_scoring_data(r), r.supplier_name) for r in raw])
            self.assertEqual(len(scoring), 2)
            self.assertTrue(all(np.isfinite(r.total_internal_score) and 0 <= r.total_internal_score <= 100
                                for r in scoring))
            selection = FinalSelectionModule()
            proposals = {r.supplier_id: reader.convert_to_supplier_proposal(r, ChemicalType.COAGULANT, 1000)
                         for r in raw}
            results = selection.finalize_selection(reader.get_all_suppliers(), proposals, {}, {}, {}, {},
                [ChemicalRequirement(ChemicalType.COAGULANT, 1000, 'synthetic standard', '普通', 5000000)],
                WastewaterType.DOMESTIC, {r.supplier_id: r for r in scoring})
            self.assertEqual({r.supplier_id for r in results}, {'A', 'B'})
            expected = sorted(scoring, key=lambda r: r.total_internal_score, reverse=True)
            self.assertEqual([r.supplier_id for r in results], [r.supplier_id for r in expected])
            xlsx, docx, txt = [Path(temp) / ('report.' + ext) for ext in ('xlsx', 'docx', 'txt')]
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertTrue(selection.export_selection_results_excel(results, WastewaterType.DOMESTIC, str(xlsx)))
                self.assertTrue(selection.export_selection_results_word(results, WastewaterType.DOMESTIC, str(docx)))
                self.assertTrue(selection.export_selection_results(results, str(txt)))
            book = load_workbook(xlsx)
            self.assertEqual(book.sheetnames, ['最终选择结果', '合同条款'])
            self.assertEqual(book.worksheets[0]['B5'].value, results[0].supplier_id)
            book.close()
            document = Document(docx)
            self.assertEqual(document.tables[0].rows[1].cells[1].text, results[0].supplier_id)
            self.assertIn(results[0].supplier_name, txt.read_text(encoding='utf-8'))


class ReleaseConfigurationTests(unittest.TestCase):
    def test_final_gate_requires_independent_review_and_successful_tests(self):
        # Synthetic records live outside the real workflow run.
        with tempfile.TemporaryDirectory() as temp:
            handoffs = Path(temp) / 'handoffs'
            handoffs.mkdir()
            base = dict(schema_version=1, task_id='synthetic-release', attempt=1,
                        started_at='2026-10-05T00:00:00Z', finished_at='2026-10-05T00:00:01Z',
                        owned_paths=[], summary='Synthetic gate regression', evidence=[])
            developer = dict(base, role='developer', actor_id='developer', status='implemented',
                             owned_paths=['main.py'], changed_files=['main.py'])
            reviewer = dict(base, role='reviewer', actor_id='reviewer', status='approved',
                            verdict='approved', findings=[], started_at='2026-10-05T00:00:02Z',
                            finished_at='2026-10-05T00:00:03Z')
            tester = dict(reviewer, role='tester', actor_id='tester', status='passed',
                          commands=[dict(command='synthetic test', exit_code=0)], gaps=[])
            for case in ('approved', 'same_actor', 'failed_test', 'stale_verification'):
                records = [dict(developer), dict(reviewer), dict(tester)]
                if case == 'same_actor':
                    records[2]['actor_id'] = 'developer'
                elif case == 'failed_test':
                    records[2]['commands'] = [dict(command='synthetic test', exit_code=1)]
                elif case == 'stale_verification':
                    records[0]['finished_at'] = '2026-10-05T00:00:04Z'
                for record in records:
                    (handoffs / (record['role'] + '.json')).write_text(json.dumps(record), encoding='utf-8')
                result = subprocess.run([sys.executable, str(ROOT / '.ai-workflow/scripts/workflow_gate.py'),
                    'final', temp], capture_output=True, timeout=20)
                with self.subTest(case=case):
                    self.assertEqual(result.returncode, 0 if case == 'approved' else 2, result.stdout)
                    self.assertEqual(json.loads(result.stdout)['success'], case == 'approved')

    def test_core_exports_dependencies_documented_and_ci_runs_suite(self):
        requirements = (ROOT / 'requirements.txt').read_text(encoding='utf-8')
        for dependency in ('openpyxl', 'python-docx', 'pandas', 'numpy', 'scikit-learn'):
            self.assertIn(dependency, requirements)
        readme = (ROOT / 'README.md').read_text(encoding='utf-8')
        for command in ('python main.py --help', 'python main.py --hdp --help',
                        'python -m unittest discover -s tests -v'):
            self.assertIn(command, readme)
        workflow = (ROOT / '.github/workflows/tests.yml').read_text(encoding='utf-8')
        self.assertIn('python -m pip install -r requirements.txt', workflow)
        self.assertIn('python -m unittest discover -s tests -v', workflow)

    def test_sensitive_and_generated_paths_are_git_ignored(self):
        paths = ['config/secret_key', 'config/secret_key.local', '.env', '.env.production',
                 'data/raw/private.xlsx', 'data/models/model.pt', 'backup/private.xlsx',
                 '.vscode/settings.json', 'reports/report.docx', 'selection_results.xlsx']
        result = subprocess.run(['git', 'check-ignore', '--no-index', '-z', '--stdin'],
            input=('\0'.join(paths) + '\0').encode(), cwd=ROOT, capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(set(result.stdout.decode().rstrip('\0').split('\0')), set(paths))

    def test_hdp_model_sources_publish_while_root_model_artifacts_are_ignored(self):
        sources = ['HDP_DS/models/__init__.py', 'HDP_DS/models/hdp_model.py',
                   'HDP_DS/models/static_encoder.py', 'HDP_DS/models/tft_encoder.py']
        artifacts = ['models/local_model.bin', 'data/models/local_model.bin']
        result = subprocess.run(['git', 'check-ignore', '--no-index', '-z', '--stdin'],
            input=('\0'.join(sources + artifacts) + '\0').encode(),
            cwd=ROOT, capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        ignored = set(result.stdout.decode().rstrip('\0').split('\0'))
        self.assertEqual(ignored, set(artifacts))
        self.assertTrue((ROOT / sources[1]).is_file(), 'Required HDP model source is missing')


if __name__ == '__main__':
    unittest.main()
